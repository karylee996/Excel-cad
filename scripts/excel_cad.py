from __future__ import annotations

import argparse
import csv
import math
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Optional

import ezdxf
from ezdxf import bbox
from ezdxf.addons import Importer
from ezdxf.enums import TextEntityAlignment
from openpyxl import load_workbook

HEITI_STYLE = "HEITI"
HEITI_FONT = "simhei.ttf"
TEXT_COLOR = 7
FRAME_COLOR = 8
MISSING_COLOR = 1
CODE_RE = re.compile(r"([A-Z]{2}\d{4})")
ZONE_HEADERS = {"区域", "区域号", "分区", "area", "zone"}
CODE_HEADERS = {"编号", "家具编号", "产品编号", "code", "model", "型号"}
QTY_HEADERS = {"数量", "qty", "quantity", "数目"}


@dataclass
class ScheduleItem:
    zone: str
    code: str
    qty: int

    @property
    def key(self) -> Optional[str]:
        return match_key(self.code)


@dataclass
class LibraryItem:
    block_name: str
    key: str
    brand: str = ""


def norm(v) -> str:
    return "" if v is None else str(v).strip()


def normalize_code(code: str) -> str:
    return re.sub(r"[^A-Z0-9]", "", norm(code).upper())


def match_key(code: str) -> Optional[str]:
    s = normalize_code(code)
    m = CODE_RE.search(s)
    if m:
        return m.group(1)
    if len(s) >= 6 and any(c.isalpha() for c in s[:6]) and any(c.isdigit() for c in s[:6]):
        return s[:6]
    return None


def normalize_zone(v) -> str:
    s = norm(v)
    m = re.search(r"\d+", s)
    if m:
        return f"{int(m.group()):02d}"
    return s.replace("区", "")


def safe_int(v, default=1) -> int:
    try:
        return max(1, int(float(v))) if norm(v) else default
    except Exception:
        return default


def detect_groups(ws):
    for row in range(1, min(ws.max_row, 12) + 1):
        headers = {c: norm(ws.cell(row, c).value).lower().replace(" ", "") for c in range(1, ws.max_column + 1)}
        groups = []
        for zc, h in headers.items():
            if h not in ZONE_HEADERS:
                continue
            cc = qc = None
            for c in range(zc + 1, min(ws.max_column, zc + 6) + 1):
                if cc is None and headers[c] in CODE_HEADERS:
                    cc = c
                elif qc is None and headers[c] in QTY_HEADERS:
                    qc = c
            if cc and qc:
                groups.append((zc, cc, qc, row))
        if groups:
            return groups
    return []


def parse_excel(path: Path, sheet: Optional[str] = None):
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.active
    groups = detect_groups(ws)
    if not groups:
        raise ValueError("未识别到 Excel 中的 区域/编号/数量 列组")
    out = []
    for zc, cc, qc, header_row in groups:
        current_zone = ""
        for r in range(header_row + 1, ws.max_row + 1):
            zv = norm(ws.cell(r, zc).value)
            cv = norm(ws.cell(r, cc).value)
            if zv:
                current_zone = normalize_zone(zv)
            if not cv:
                continue
            if cv.lower().replace(" ", "") in CODE_HEADERS:
                continue
            zone = normalize_zone(zv) if zv else current_zone
            if zone:
                out.append(ScheduleItem(zone, cv, safe_int(ws.cell(r, qc).value)))
    return out


def aggregate(items):
    grouped = defaultdict(int)
    display = {}
    for it in items:
        k = (it.zone, normalize_code(it.code) or it.code.upper())
        grouped[k] += it.qty
        display[k] = it.code
    out = defaultdict(list)
    for (zone, nk), qty in grouped.items():
        out[zone].append(ScheduleItem(zone, display[(zone, nk)], qty))
    for z in out:
        out[z].sort(key=lambda x: (x.key or "ZZZZZZ", normalize_code(x.code)))
    return out


def ensure_heiti(doc):
    if HEITI_STYLE not in doc.styles:
        doc.styles.add(HEITI_STYLE, font=HEITI_FONT)
    else:
        doc.styles.get(HEITI_STYLE).dxf.font = HEITI_FONT


def force_heiti(doc):
    ensure_heiti(doc)
    for layout in doc.layouts:
        for e in layout:
            if e.dxftype() in {"TEXT", "MTEXT", "ATTRIB", "ATTDEF"}:
                try:
                    e.dxf.style = HEITI_STYLE
                except Exception:
                    pass
    for block in doc.blocks:
        for e in block:
            if e.dxftype() in {"TEXT", "MTEXT", "ATTRIB", "ATTDEF"}:
                try:
                    e.dxf.style = HEITI_STYLE
                except Exception:
                    pass


def entity_text(e):
    if e.dxftype() == "TEXT":
        return e.dxf.text.strip()
    if e.dxftype() == "MTEXT":
        try:
            return e.plain_text().strip()
        except Exception:
            return e.text.strip()
    return ""


def scan_library(doc):
    msp = doc.modelspace()
    brand_texts = []
    for typ in ("TEXT", "MTEXT"):
        for e in msp.query(typ):
            if e.dxf.layer.upper() == "BRAND":
                brand_texts.append((e.dxf.insert.x, e.dxf.insert.y, entity_text(e)))
    found = defaultdict(list)
    for ins in msp.query("INSERT"):
        key = match_key(ins.dxf.name)
        if not key:
            continue
        brand = ""
        if brand_texts:
            x, y = ins.dxf.insert.x, ins.dxf.insert.y
            best = min(brand_texts, key=lambda q: (q[0]-x)**2 + (q[1]-y)**2)
            if math.hypot(best[0]-x, best[1]-y) < 1200:
                brand = best[2]
        found[key].append(LibraryItem(ins.dxf.name, key, brand))
    return found


def common_prefix(a, b):
    n = 0
    for x, y in zip(a.upper(), b.upper()):
        if x != y:
            break
        n += 1
    return n


def choose_library(code, library):
    key = match_key(code)
    if not key or key not in library:
        return None
    tc = norm(code).upper().replace(" ", "")
    return sorted(library[key], key=lambda it: (
        0 if it.block_name.upper() == tc else 1,
        -common_prefix(tc, it.block_name.upper()),
        abs(len(it.block_name) - len(tc)),
        it.block_name,
    ))[0]


def add_text(msp, text, xy, height, color=TEXT_COLOR, align=TextEntityAlignment.MIDDLE_CENTER):
    t = msp.add_text(text, dxfattribs={"height": height, "style": HEITI_STYLE, "color": color})
    t.set_placement(xy, align=align)
    return t


def find_zone_positions(msp):
    result = {}
    for typ in ("TEXT", "MTEXT"):
        for e in msp.query(typ):
            txt = entity_text(e).strip().replace(" ", "")
            m = re.fullmatch(r"(0?[1-9]|[12]\d)(区)?", txt)
            if m and 1 <= int(m.group(1)) <= 99:
                result[f"{int(m.group(1)):02d}"] = (e.dxf.insert.x, e.dxf.insert.y)
    return result


def panel_dims(n, side, cell_w, cell_h, title_h, pad):
    n = max(1, n)
    if side in ("T", "B"):
        cols = min(8, max(3, math.ceil(math.sqrt(n * 1.6))))
    else:
        cols = min(6, max(3, math.ceil(math.sqrt(n * 1.2))))
    rows = math.ceil(n / cols)
    return cols * cell_w + 2 * pad, title_h + rows * cell_h + 2 * pad, cols


def pack_panels(area_pos, zones, base_ext, cell_w, cell_h, title_h, pad, margin, gap):
    minx, miny, maxx, maxy = base_ext
    sides = defaultdict(list)
    for a in zones:
        x, y = area_pos.get(a, ((minx+maxx)/2, (miny+maxy)/2))
        d = {"L": x-minx, "R": maxx-x, "T": maxy-y, "B": y-miny}
        sides[min(d, key=d.get)].append(a)
    sides["T"].sort(key=lambda a: area_pos.get(a, (0, 0))[0])
    sides["B"].sort(key=lambda a: area_pos.get(a, (0, 0))[0])
    sides["L"].sort(key=lambda a: area_pos.get(a, (0, 0))[1], reverse=True)
    sides["R"].sort(key=lambda a: area_pos.get(a, (0, 0))[1], reverse=True)

    dims = {a: panel_dims(len(zones[a]), s, cell_w, cell_h, title_h, pad) for s, arr in sides.items() for a in arr}
    placed = {}

    def horizontal(arr, top):
        rows = []
        for a in arr:
            w, h, _ = dims[a]
            ax = area_pos.get(a, ((minx+maxx)/2, 0))[0]
            chosen = None
            for ri, row in enumerate(rows):
                last = row[-1][2] if row else -1e99
                x = max(ax-w/2, last+gap)
                if x+w <= maxx+18000:
                    chosen = (ri, x)
                    break
            if chosen is None:
                ri = len(rows); rows.append([]); chosen = (ri, max(minx-12000, ax-w/2))
            ri, x = chosen
            rows[ri].append((a, x, x+w, h))
        for ri, row in enumerate(rows):
            if not row:
                continue
            shift = (minx+maxx)/2 - (min(r[1] for r in row)+max(r[2] for r in row))/2
            outward = sum(max(rr[3] for rr in rows[k])+gap for k in range(ri))
            for a, x1, x2, h in row:
                w = x2-x1
                y = maxy+margin+outward if top else miny-margin-h-outward
                placed[a] = (x1+shift, y, w, h, "T" if top else "B")

    def vertical(arr, left):
        cols = []
        for a in arr:
            w, h, _ = dims[a]
            ay = area_pos.get(a, (0, (miny+maxy)/2))[1]
            chosen = None
            for ci, col in enumerate(cols):
                last = col[-1][2] if col else 1e99
                yt = min(ay+h/2, last-gap)
                if yt-h >= miny-15000:
                    chosen = (ci, yt)
                    break
            if chosen is None:
                ci = len(cols); cols.append([]); chosen = (ci, min(maxy+12000, ay+h/2))
            ci, yt = chosen
            cols[ci].append((a, yt, yt-h, w, h))
        for ci, col in enumerate(cols):
            shift = (miny+maxy)/2 - (max(r[1] for r in col)+min(r[2] for r in col))/2
            outward = sum(max(rr[3] for rr in cols[k])+gap for k in range(ci))
            for a, yt, yb, w, h in col:
                x = minx-margin-w-outward if left else maxx+margin+outward
                placed[a] = (x, yb+shift, w, h, "L" if left else "R")

    horizontal(sides["T"], True)
    horizontal(sides["B"], False)
    vertical(sides["L"], True)
    vertical(sides["R"], False)
    return placed, dims, sides


def build(excel_path: Path, base_path: Path, library_path: Path, output: Path, report: Path, sheet=None):
    schedule = parse_excel(excel_path, sheet)
    zones = aggregate(schedule)
    base = ezdxf.readfile(base_path)
    lib = ezdxf.readfile(library_path)
    force_heiti(base)
    msp = base.modelspace()
    library = scan_library(lib)

    resolved = []
    needed = set()
    for zone, items in zones.items():
        for item in items:
            li = choose_library(item.code, library)
            resolved.append((zone, item, li))
            if li:
                needed.add(li.block_name)

    importer = Importer(lib, base)
    for block in sorted(needed):
        try:
            importer.import_block(block)
        except Exception:
            pass
    importer.finalize()

    sizes, centers = {}, {}
    for b in needed:
        try:
            ex = bbox.extents(base.blocks.get(b))
            sizes[b] = (max(ex.size.x, 1), max(ex.size.y, 1))
            centers[b] = ((ex.extmin.x+ex.extmax.x)/2, (ex.extmin.y+ex.extmax.y)/2)
        except Exception:
            sizes[b] = (1800, 1200); centers[b] = (0, 0)

    ext = bbox.extents(msp)
    base_ext = (ext.extmin.x, ext.extmin.y, ext.extmax.x, ext.extmax.y)
    area_pos = find_zone_positions(msp)

    CELL_W, CELL_H, TITLE_H, PAD, MARGIN, GAP = 2700.0, 1950.0, 900.0, 350.0, 2200.0, 900.0
    placed, dims, sides = pack_panels(area_pos, zones, base_ext, CELL_W, CELL_H, TITLE_H, PAD, MARGIN, GAP)

    for zone in sorted(zones, key=lambda s: int(re.search(r"\d+", s).group()) if re.search(r"\d+", s) else 999):
        x, y, w, h, side = placed[zone]
        msp.add_lwpolyline([(x,y),(x+w,y),(x+w,y+h),(x,y+h)], close=True, dxfattribs={"color":FRAME_COLOR})
        add_text(msp, f"{zone}区", (x+PAD, y+h-500), 420, align=TextEntityAlignment.LEFT)
        cols = dims[zone][2]
        content_top = y+h-TITLE_H-PAD
        for idx, item in enumerate(zones[zone]):
            c, row = idx % cols, idx // cols
            cx = x+PAD+c*CELL_W+CELL_W/2
            cy = content_top-row*CELL_H-CELL_H/2
            li = choose_library(item.code, library)
            if li:
                bw, bh = sizes[li.block_name]
                s = min(1.0, 1950.0/bw, 1150.0/bh)
                lcx, lcy = centers[li.block_name]
                target_y = cy+80
                msp.add_blockref(li.block_name, (cx-lcx*s, target_y-lcy*s), dxfattribs={"xscale":s,"yscale":s,"zscale":s})
                add_text(msp, item.code, (cx,target_y+105), 170)
                if li.brand:
                    add_text(msp, li.brand, (cx,target_y-105), 145)
                if item.qty != 1:
                    add_text(msp, f"x{item.qty}", (cx+CELL_W/2-180, cy-CELL_H/2+180), 145, align=TextEntityAlignment.MIDDLE_RIGHT)
            else:
                add_text(msp, item.code, (cx,cy+30), 220, color=MISSING_COLOR)
                if item.qty != 1:
                    add_text(msp, f"x{item.qty}", (cx,cy-230), 145, color=MISSING_COLOR)

    force_heiti(base)
    base.saveas(output)
    ezdxf.readfile(output)

    with report.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.writer(f)
        w.writerow(["区域", "原始编码", "前6位匹配键", "数量", "家具库块", "品牌", "状态"])
        for zone, item, li in resolved:
            w.writerow([f"{zone}区", item.code, item.key or "", item.qty, li.block_name if li else "", li.brand if li else "", "matched" if li else "missing"])

    return {"matched": sum(1 for _,_,li in resolved if li), "missing": sum(1 for _,_,li in resolved if not li), "sides": dict(sides)}


def main():
    p = argparse.ArgumentParser(description="Excel + CAD base + furniture library -> nearby zone furniture index DXF")
    p.add_argument("--excel", required=True, type=Path)
    p.add_argument("--base", required=True, type=Path)
    p.add_argument("--library", required=True, type=Path)
    p.add_argument("--output", required=True, type=Path)
    p.add_argument("--report", type=Path)
    p.add_argument("--sheet")
    args = p.parse_args()
    report = args.report or args.output.with_name(args.output.stem + "_match_report.csv")
    print(build(args.excel, args.base, args.library, args.output, report, args.sheet))


if __name__ == "__main__":
    main()
